"""Firestore state를 주간 로컬 백업으로 내보낸다."""

from __future__ import annotations

import argparse
import getpass
import html
import json
import re
import shutil
import sys
import traceback
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE


# src/data/cache.ts CACHE_KEYS와 같게 유지
KEYS = (
    "records", "completed", "meta", "study", "studyFiles", "events",
    "rdda", "fabricAnalysis", "ts", "orgMembers", "materials",
    "materialsManual", "materialDiagnostics", "fabricOverrides",
    "fabricEvents", "chemical", "chemicalManual", "chemicalLinks",
    "requests",
)

LABELS = {
    "records": "DD MASTER",
    "fabricOverrides": "창고 상태",
    "fabricEvents": "창고 이력",
    "completed": "샘플대장",
    "ts": "TROUBLE SHOOTING",
    "requests": "FABRIC REQUEST",
}

SHEETS = (
    ("DD", "records"),
    ("창고상태", "fabricOverrides"),
    ("창고이력", "fabricEvents"),
    ("샘플대장", "completed"),
    ("TS", "ts"),
    ("REQUEST", "requests"),
)

BACKUP_NAME_RE = re.compile(r"^FabricRnD_주간백업_(\d{8})(\.zip)?$")
CELL_TEXT_LIMIT = 32_000
CREDENTIAL_TARGET = "FabricRnD-weekly-backup"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Fabric R&D 주간 로컬 백업")
    parser.add_argument(
        "--config", type=Path,
        default=Path.home() / "fabric-backup" / "config.json",
    )
    parser.add_argument("--no-mail", action="store_true")
    parser.add_argument("--save-login", action="store_true")
    return parser.parse_args()


def load_config(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        config = json.load(handle)
    required = (
        "project_id", "api_key", "login_email", "output_dir", "mail_to",
        "keep_weeks", "attach_limit_mb",
    )
    missing = [key for key in required if key not in config]
    if missing:
        raise ValueError(f"설정 항목 누락: {', '.join(missing)}")
    if int(config["keep_weeks"]) < 0 or float(config["attach_limit_mb"]) < 0:
        raise ValueError("keep_weeks와 attach_limit_mb는 0 이상이어야 합니다.")
    return config


def load_password() -> str:
    import win32cred

    try:
        credential = win32cred.CredRead(
            CREDENTIAL_TARGET, win32cred.CRED_TYPE_GENERIC,
        )
    except Exception:
        raise RuntimeError(
            "저장된 로그인이 없습니다. --save-login을 먼저 실행하세요."
        ) from None
    password = credential.get("CredentialBlob")
    if isinstance(password, bytes):
        return password.decode("utf-16-le")
    if isinstance(password, str):
        return password
    raise RuntimeError("저장된 로그인 형식이 올바르지 않습니다.")


def sign_in(api_key: str, email: str, password: str) -> str:
    query = urllib.parse.urlencode({"key": api_key})
    url = (
        "https://identitytoolkit.googleapis.com/v1/"
        f"accounts:signInWithPassword?{query}"
    )
    body = json.dumps({
        "email": email, "password": password, "returnSecureToken": True,
    }).encode("utf-8")
    request = urllib.request.Request(
        url, data=body,
        headers={"Content-Type": "application/json"}, method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=60) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        try:
            error_payload = json.loads(exc.read().decode("utf-8"))
            message = error_payload.get("error", {}).get("message", f"HTTP {exc.code}")
        except (UnicodeDecodeError, json.JSONDecodeError):
            message = f"HTTP {exc.code}"
        raise RuntimeError(f"로그인 실패: {message}") from None
    except urllib.error.URLError:
        raise RuntimeError("로그인 실패: 네트워크 오류") from None
    token = payload.get("idToken")
    if not isinstance(token, str):
        raise RuntimeError("로그인 실패: ID 토큰이 없습니다.")
    return token


def save_login(config: dict[str, Any]) -> int:
    import win32cred

    print(f"로그인 이메일: {config['login_email']}")
    password = getpass.getpass("앱 비밀번호: ")
    try:
        sign_in(config["api_key"], config["login_email"], password)
    except RuntimeError as exc:
        print(str(exc))
        return 1
    win32cred.CredWrite({
        "Type": win32cred.CRED_TYPE_GENERIC,
        "TargetName": CREDENTIAL_TARGET,
        "UserName": config["login_email"],
        "CredentialBlob": password,
        "Persist": win32cred.CRED_PERSIST_LOCAL_MACHINE,
    }, 0)
    print("저장했습니다")
    return 0


def field(fields: dict[str, Any], name: str) -> Any:
    value = fields.get(name, {})
    if "integerValue" in value:
        try:
            return int(value["integerValue"])
        except (TypeError, ValueError):
            return None
    if "stringValue" in value:
        return value["stringValue"]
    if "timestampValue" in value:
        return value["timestampValue"]
    if "doubleValue" in value:
        try:
            return float(value["doubleValue"])
        except (TypeError, ValueError):
            return None
    if "nullValue" in value:
        return None
    return None


class FirestoreRest:
    def __init__(self, project_id: str, id_token: str) -> None:
        self.base_path = f"projects/{project_id}/databases/(default)/documents"
        self.base_url = f"https://firestore.googleapis.com/v1/{self.base_path}"
        self.headers = {"Authorization": f"Bearer {id_token}"}

    @staticmethod
    def _read_error(exc: urllib.error.HTTPError, doc_path: str) -> RuntimeError:
        message = f"Firestore 읽기 실패 {exc.code}: {doc_path}"
        if exc.code == 403:
            message += " (앱 계정 승인 여부를 확인하세요)"
        return RuntimeError(message)

    def get(self, doc_path: str) -> dict[str, Any] | None:
        encoded_path = urllib.parse.quote(doc_path, safe="/")
        request = urllib.request.Request(
            f"{self.base_url}/{encoded_path}", headers=self.headers, method="GET",
        )
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                return json.loads(response.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            if exc.code == 404:
                return None
            raise self._read_error(exc, doc_path) from None
        except urllib.error.URLError:
            raise RuntimeError(f"Firestore 읽기 실패: {doc_path}") from None

    def batch_get(self, doc_paths: list[str]) -> dict[str, dict[str, Any] | None]:
        results: dict[str, dict[str, Any] | None] = {
            path: None for path in doc_paths
        }
        for start in range(0, len(doc_paths), 100):
            batch = doc_paths[start:start + 100]
            documents = [f"{self.base_path}/{path}" for path in batch]
            body = json.dumps({"documents": documents}).encode("utf-8")
            request = urllib.request.Request(
                f"{self.base_url}:batchGet", data=body,
                headers={**self.headers, "Content-Type": "application/json"},
                method="POST",
            )
            try:
                with urllib.request.urlopen(request, timeout=60) as response:
                    payload = json.loads(response.read().decode("utf-8"))
            except urllib.error.HTTPError as exc:
                raise self._read_error(exc, "batchGet") from None
            except urllib.error.URLError:
                raise RuntimeError("Firestore 읽기 실패: batchGet") from None
            if not isinstance(payload, list):
                raise RuntimeError("Firestore 읽기 실패: batchGet 응답 형식")
            prefix = f"{self.base_path}/"
            for item in payload:
                if "found" in item:
                    found = item["found"]
                    name = found.get("name", "")
                    if name.startswith(prefix):
                        results[name[len(prefix):]] = found
                elif "missing" in item:
                    missing = item["missing"]
                    if missing.startswith(prefix):
                        results[missing[len(prefix):]] = None
        return results


def item_count(value: Any) -> int:
    if value is None:
        return 0
    if isinstance(value, (list, dict)):
        return len(value)
    return 1


def read_state(store: FirestoreRest) -> tuple[dict[str, Any], dict[str, dict[str, Any]]]:
    data: dict[str, Any] = {}
    summary: dict[str, dict[str, Any]] = {}
    for key in KEYS:
        meta = store.get(f"state/{key}")
        if meta is None:
            data[key] = None
            summary[key] = {
                "count": 0, "updatedAt": None, "updatedBy": None, "status": "없음",
            }
            continue

        meta_fields = meta.get("fields", {})
        n = field(meta_fields, "n")
        if not isinstance(n, int) or isinstance(n, bool) or n < 0:
            raise ValueError(f"{key}: 잘못된 조각 수")
        paths = [f"state/{key}__{index}" for index in range(n)]
        snapshots = store.batch_get(paths)
        chunks: list[str] = []
        for index in range(n):
            document_path = f"state/{key}__{index}"
            snapshot = snapshots.get(document_path)
            if snapshot is None:
                raise ValueError(f"{key}: 조각 {index} 누락")
            chunk = field(snapshot.get("fields", {}), "c")
            if not isinstance(chunk, str):
                raise ValueError(f"{key}: 조각 {index}의 c가 문자열이 아님")
            chunks.append(chunk)
        try:
            value = json.loads("".join(chunks))
        except json.JSONDecodeError as exc:
            raise ValueError(f"{key}: JSON 파싱 실패") from exc
        data[key] = value
        summary[key] = {
            "count": item_count(value),
            "updatedAt": field(meta_fields, "updatedAt"),
            "updatedBy": field(meta_fields, "updatedBy"),
            "status": "정상",
        }
    return data, summary


def json_dump(path: Path, value: Any) -> None:
    with path.open("w", encoding="utf-8") as handle:
        json.dump(value, handle, ensure_ascii=False, indent=2)
        handle.write("\n")


def cell_value(value: Any) -> str | int | float | bool:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False)
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    return text[:CELL_TEXT_LIMIT] + "…(생략)" if len(text) > CELL_TEXT_LIMIT else text


def flatten_row(
    value: Any, prefix: str = "", result: dict[str, Any] | None = None,
) -> dict[str, Any]:
    flattened = {} if result is None else result
    if not isinstance(value, dict):
        if prefix:
            flattened[prefix] = cell_value(value)
        return flattened
    for key, child in value.items():
        path = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(child, dict):
            flatten_row(child, path, flattened)
        else:
            flattened[path] = cell_value(child)
    return flattened


def build_workbook(
    path: Path, exported_at: str, data: dict[str, Any],
    summary: dict[str, dict[str, Any]],
) -> None:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "안내"
    guide.append(["내보낸 시각", exported_at])
    guide.append(["안내", "사람이 읽는 백업입니다. 복원에는 backup.json을 쓰세요."])
    guide.append([])
    guide.append(["키", "이름", "건수", "마지막 수정 시각", "수정자"])
    for key in KEYS:
        item = summary[key]
        guide.append([
            key, LABELS.get(key, key), item["count"],
            cell_value(item["updatedAt"]), cell_value(item["updatedBy"]),
        ])
    for cell in guide[4]:
        cell.font = Font(bold=True)
    guide.freeze_panes = "A2"

    for name, key in SHEETS:
        worksheet = workbook.create_sheet(name)
        values = data[key]
        if not isinstance(values, list):
            continue
        rows = [flatten_row(value) for value in values]
        headers = list(dict.fromkeys(column for row in rows for column in row))
        if not headers:
            continue
        worksheet.append([cell_value(header) for header in headers])
        for row in rows:
            worksheet.append([row.get(header, "") for header in headers])
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
        worksheet.freeze_panes = "A2"
        for index, header in enumerate(headers, 1):
            column = worksheet.cell(1, index).column_letter
            worksheet.column_dimensions[column].width = min(48, max(12, len(header) + 2))
    workbook.save(path)


def find_previous_summary(output_dir: Path, current_dir: Path) -> dict[str, Any] | None:
    candidates: list[tuple[str, Path]] = []
    for path in output_dir.iterdir():
        match = BACKUP_NAME_RE.fullmatch(path.name)
        if path.is_dir() and match and not match.group(2) and path != current_dir:
            candidates.append((match.group(1), path))
    for _, folder in sorted(candidates, reverse=True):
        summary_path = folder / "summary.json"
        if summary_path.is_file():
            with summary_path.open("r", encoding="utf-8") as handle:
                return json.load(handle)
    return None


def weekly_changes(
    summary: dict[str, dict[str, Any]], previous: dict[str, Any] | None,
) -> dict[str, int | str]:
    if previous is None:
        return {key: "" for key in KEYS}
    changes: dict[str, int | str] = {}
    for key in KEYS:
        old = previous.get(key, {}).get("count")
        changes[key] = summary[key]["count"] - old if isinstance(old, int) else ""
    return changes


def create_zip(zip_path: Path, backup_dir: Path, files: list[Path]) -> None:
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, arcname=path.relative_to(backup_dir))


def prune_backups(output_dir: Path, keep_weeks: int, today: date) -> None:
    cutoff = today - timedelta(days=keep_weeks * 7)
    for path in output_dir.iterdir():
        match = BACKUP_NAME_RE.fullmatch(path.name)
        if not match:
            continue
        try:
            item_date = datetime.strptime(match.group(1), "%Y%m%d").date()
        except ValueError:
            continue
        if item_date >= cutoff:
            continue
        if path.is_dir() and not match.group(2):
            shutil.rmtree(path)
        elif path.is_file() and match.group(2):
            path.unlink()


def format_delta(value: int | str) -> str:
    return "" if value == "" else f"{value:+d}"


def mail_table(summary: dict[str, dict[str, Any]], changes: dict[str, int | str]) -> str:
    rows = []
    for key in KEYS:
        item = summary[key]
        delta = changes[key]
        previous = "" if delta == "" else item["count"] - delta
        cells = (
            LABELS.get(key, key), item["count"], previous,
            format_delta(delta), item["updatedAt"] or "",
        )
        rows.append("<tr>" + "".join(f"<td>{html.escape(str(cell))}</td>" for cell in cells) + "</tr>")
    return (
        "<table border='1' cellspacing='0' cellpadding='4'>"
        "<tr><th>이름</th><th>건수</th><th>지난주</th><th>증감</th><th>마지막 수정 시각</th></tr>"
        + "".join(rows) + "</table>"
    )


def send_mail(to: str, subject: str, body: str, attachment: Path | None = None) -> None:
    import win32com.client

    outlook = win32com.client.Dispatch("Outlook.Application")
    message = outlook.CreateItem(0)
    message.To = to
    message.Subject = subject
    message.HTMLBody = body
    if attachment is not None:
        message.Attachments.Add(str(attachment))
    message.Send()


def append_log(output_dir: Path, outcome: str, detail: str) -> None:
    timestamp = datetime.now().astimezone().isoformat(timespec="seconds")
    with (output_dir / "backup.log").open("a", encoding="utf-8") as handle:
        handle.write(f"{timestamp} {outcome} {detail}\n")


def run_backup(config: dict[str, Any]) -> dict[str, Any]:
    output_dir = Path(config["output_dir"])
    output_dir.mkdir(parents=True, exist_ok=True)
    password = load_password()
    token = sign_in(config["api_key"], config["login_email"], password)
    store = FirestoreRest(config["project_id"], token)
    data, summary = read_state(store)
    now = datetime.now().astimezone()
    stamp = now.strftime("%Y%m%d")
    backup_dir = output_dir / f"FabricRnD_주간백업_{stamp}"
    backup_dir.mkdir(parents=True, exist_ok=True)
    exported_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    backup_path = backup_dir / "backup.json"
    summary_path = backup_dir / "summary.json"
    workbook_path = backup_dir / f"FabricRnD_주간백업_{stamp}.xlsx"
    zip_path = output_dir / f"FabricRnD_주간백업_{stamp}.zip"
    json_dump(backup_path, {
        "app": "fabric-rnd", "version": 1, "exportedAt": exported_at,
        "source": "weekly-backup", "data": data,
    })
    json_dump(summary_path, summary)
    build_workbook(workbook_path, exported_at, data, summary)
    create_zip(zip_path, backup_dir, [backup_path, summary_path, workbook_path])
    previous = find_previous_summary(output_dir, backup_dir)
    changes = weekly_changes(summary, previous)
    prune_backups(output_dir, int(config["keep_weeks"]), now.date())
    return {
        "output_dir": output_dir, "backup_dir": backup_dir, "zip_path": zip_path,
        "date": now.date(), "summary": summary, "changes": changes,
    }


def main() -> int:
    args = parse_args()
    config: dict[str, Any] | None = None
    output_dir: Path | None = None
    if args.save_login:
        try:
            config = load_config(args.config)
            return save_login(config)
        except Exception as exc:
            print(f"로그인 저장 실패: {type(exc).__name__}: {exc}")
            return 1
    try:
        config = load_config(args.config)
        output_dir = Path(config["output_dir"])
        result = run_backup(config)
        counts = ", ".join(f"{key}={result['summary'][key]['count']}" for key in KEYS)
        append_log(result["output_dir"], "성공", counts)
        if not args.no_mail:
            zip_path = result["zip_path"]
            limit = float(config["attach_limit_mb"]) * 1024 * 1024
            attachment = zip_path if zip_path.stat().st_size <= limit else None
            attachment_note = "" if attachment else (
                f"<p>용량 초과로 첨부하지 않았습니다: {html.escape(str(zip_path))}</p>"
            )
            body = (
                mail_table(result["summary"], result["changes"])
                + f"<p>저장 경로: {html.escape(str(result['backup_dir']))}</p>"
                + attachment_note
            )
            try:
                send_mail(
                    config["mail_to"],
                    f"[Fabric R&D] 주간 백업 완료 {result['date'].isoformat()}",
                    body, attachment,
                )
            except Exception as exc:
                append_log(result["output_dir"], "메일실패", f"{type(exc).__name__}: {exc}")
                print("주간 백업은 완료됐지만 메일 발송에 실패했습니다.")
                return 2
        print(f"주간 백업 완료: {result['backup_dir']}")
        return 0
    except Exception as exc:
        trace_lines = traceback.format_exc().splitlines()[-15:]
        if output_dir is not None:
            try:
                output_dir.mkdir(parents=True, exist_ok=True)
                append_log(output_dir, "실패", f"{type(exc).__name__}: {exc}")
            except Exception:
                pass
        if not args.no_mail and config is not None:
            body = (
                f"<p>{html.escape(type(exc).__name__)}: {html.escape(str(exc))}</p>"
                f"<pre>{html.escape(chr(10).join(trace_lines))}</pre>"
            )
            try:
                send_mail(
                    config["mail_to"],
                    f"[Fabric R&D] 주간 백업 실패 {date.today().isoformat()}", body,
                )
            except Exception as mail_exc:
                if output_dir is not None:
                    try:
                        append_log(output_dir, "메일실패", f"{type(mail_exc).__name__}: {mail_exc}")
                    except Exception:
                        pass
                print("주간 백업과 실패 메일 발송에 실패했습니다.")
                return 2
        print(f"주간 백업 실패: {type(exc).__name__}: {exc}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
